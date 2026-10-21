import asyncio
import json

import httpx
import pytest

from worker.agents import build_agent, parse_json, run_agent
from worker.config import RunConfig, Settings
from worker.llm_client import Model, call_context, longest_gap, parse_usage
from worker.orchestrator import execute, select_by_tokens, summarise
from worker.results import RECORD_COLUMNS, append_csv, parse_call_lines
from worker.search_tool import make_search_tool
from worker.store import MemoryStore
from worker.tests import fake_model


def run(coro):
    return asyncio.run(coro)


def settings(**kw) -> Settings:
    return Settings(store_mode="memory", llm_base_url_override="http://fake/v1", **kw)


def asgi():
    return httpx.ASGITransport(app=fake_model.app)


class _Broken(httpx.AsyncByteStream):
    """Delivers some bytes, then the connection dies (what an idle-timeout cut looks like)."""

    def __init__(self, chunks: list[bytes], error: Exception | None) -> None:
        self.chunks, self.error = chunks, error

    async def __aiter__(self):
        for c in self.chunks:
            yield c
        if self.error:
            raise self.error

    async def aclose(self) -> None:
        return None


def event_chunk() -> bytes:
    """A valid opening SSE event, as a real Responses stream starts."""
    started = {**fake_model._response("m", False), "status": "in_progress", "output": []}
    return fake_model._sse({"type": "response.created", "sequence_number": 0, "response": started})


def mock_transport(status=200, chunks=(b"data: x\n\n",), error=None):
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(status, stream=_Broken(list(chunks), error))

    return httpx.MockTransport(handler)


async def call(st=None, path="direct", transport=None, attempts=None, **kw):
    recs = []

    async def rec(r):
        recs.append(r)

    m = Model(st or settings(), path, rec, max_attempts=attempts, transport=transport or asgi())
    try:
        res = await m.call("hi", max_output_tokens=100, **kw)
    finally:
        await m.aclose()
    return res, recs


@pytest.mark.parametrize("stream", [False, True])
def test_call_model_records_every_metric(stream):
    res, recs = run(call(stream=stream, label="cell-1"))
    assert res.ok and json.loads(res.text) == {"fake": True}
    r = res.record
    assert r["label"] == "cell-1" and r["http_status"] == 200 and r["outcome"] == "completed"
    assert r["ttfh_s"] is not None and r["ttfb_s"] is not None and r["total_s"] >= r["ttfb_s"]
    assert (r["prompt_tokens"], r["completion_tokens"], r["reasoning_tokens"]) == (11, 7, 5)
    assert r["request_ids"].startswith("x-request-id=") and r["client_request_id"]
    assert r["stream"] is stream and (("ttft_s" in r) is stream)
    assert len(recs) == 1
    assert set(r) - set(RECORD_COLUMNS) == set()  # every field is a CSV column


def test_silent_gap_is_measured_on_the_wire(live_fake):
    fake_model.FIRST_DELAY = 0.7
    st = Settings(store_mode="memory", llm_base_url_override=live_fake)

    async def go():
        m = Model(st, "direct", max_attempts=1)
        try:
            return await m.call("x", stream=True, max_output_tokens=50)
        finally:
            await m.aclose()

    r = run(go()).record
    assert r["ttfh_s"] < 0.5  # headers (and the first events) arrive immediately
    assert r["longest_gap_s"] >= 0.6  # then ~0.7 s of silence before the next bytes
    assert r["ttft_s"] >= 0.65  # the first token only arrives after the silence


def test_connection_cut_mid_stream_is_truncated():
    err = httpx.RemoteProtocolError("peer closed connection")
    res, recs = run(call(stream=True, transport=mock_transport(chunks=[event_chunk()], error=err)))
    r = res.record
    assert not res.ok and r["outcome"] == "truncated" and r["http_status"] == 200
    assert r["error_type"] == "RemoteProtocolError" and r["chunks"] == 1
    assert r["tail_silence_s"] is not None


def test_clean_close_without_completion_marker_is_truncated():
    res, _ = run(call(stream=True, transport=mock_transport(chunks=[event_chunk()])))
    assert res.record["outcome"] == "truncated"
    assert res.record["error_type"] == "StreamEndedWithoutCompletion"


def test_http_error_keeps_status_and_body():
    res, _ = run(
        call(transport=mock_transport(status=504, chunks=[b'{"error":"gateway timeout"}']))
    )
    r = res.record
    assert r["http_status"] == 504 and r["outcome"] == "http_error"
    assert "gateway timeout" in r["error_detail"]


def test_every_retry_attempt_is_a_record():
    st = settings(llm_max_attempts=3, llm_retry_backoff_s=0.01)
    res, recs = run(call(st=st, transport=mock_transport(status=503), attempts=3))
    assert [r["attempt"] for r in recs] == [1, 2, 3]
    assert all(r["http_status"] == 503 for r in recs) and not res.ok


def test_matrix_calls_never_retry():
    st = settings(llm_max_attempts=3)
    _, recs = run(call(st=st, transport=mock_transport(status=503), attempts=1))
    assert len(recs) == 1


def test_parse_usage_takes_last_object_from_sse_tail():
    tail = b'data: {"response":{"usage":null}}\n\ndata: {"response":{"usage":{"input_tokens":5,"output_tokens":9,"output_tokens_details":{"reasoning_tokens":4}}}}\n\n'
    assert parse_usage(tail) == {"prompt_tokens": 5, "completion_tokens": 9, "reasoning_tokens": 4}
    assert parse_usage(b"no usage here") == {}
    assert longest_gap([0.0, 1.0, 301.0]) == 300.0


def test_results_csv_roundtrip(tmp_path):
    p = tmp_path / "runs.csv"
    assert append_csv(p, {"label": "a", "total_s": 1.5}) and append_csv(p, {"label": "b"})
    lines = p.read_text().splitlines()
    assert lines[0].split(",") == RECORD_COLUMNS and len(lines) == 3
    assert parse_call_lines(["noise", 'x CALL {"label": "z"}'])[0]["label"] == "z"
    assert append_csv(tmp_path / "ro" / "x.csv", {}) is True


@pytest.mark.parametrize("stream", [False, True])
def test_agent_framework_agent_with_tool_loop(stream):
    searched = []

    async def search(query, domain, top):
        searched.append(query)
        return "[]"

    async def go():
        recs = []

        async def rec(r):
            recs.append(r)

        m = Model(settings(), "direct", rec, transport=asgi())
        tool = make_search_tool(search, rec, 3)
        with call_context(run_id="r", agent="network", path="direct", stream=stream):
            agent = build_agent(
                m.agent_client(), "network", "be brief", tools=[tool], max_output_tokens=100
            )
            out = await run_agent(agent, "go", stream=stream)
        await m.aclose()
        return out, recs

    out, recs = run(go())
    assert out == {"fake": True} and searched == ["vlan"]
    assert [r["kind"] for r in recs] == ["llm", "tool", "llm"]
    assert [r["turn"] for r in recs if r["kind"] == "llm"] == [1, 2]
    assert all(r["agent"] == "network" for r in recs)


def test_parse_json_variants():
    assert parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json('Here you go: {"a": {"b": 2}} done') == {"a": {"b": 2}}
    with pytest.raises(ValueError):
        parse_json("no json")


def test_select_by_tokens_and_summarise():
    chunks = [
        {"id": f"c{i}", "content": "x" * 400, "filename": "f", "page_number": i, "section_path": "s",
         "chunk_type": "prose"}
        for i in range(100)
    ]  # fmt: skip
    assert 0 < len(select_by_tokens(chunks, 1000)) < 100
    calls = [
        {"kind": "llm", "outcome": "completed", "ttfb_s": 1.0, "longest_gap_s": 5.0, "total_s": 9.0,
         "prompt_tokens": 10, "completion_tokens": 2, "reasoning_tokens": 1},
        {"kind": "llm", "outcome": "truncated", "ttfb_s": 3.0, "total_s": 4.0},
        {"kind": "tool", "total_s": 0.1},
    ]  # fmt: skip
    s = summarise(calls, 12.3)
    assert s["llm_calls"] == 2 and s["failed_calls"] == 1 and s["max_longest_gap_s"] == 5.0


@pytest.mark.parametrize("stream", [False, True])
def test_pipeline_end_to_end(stream):
    """Planner (with ai_search) decides the sections; only assigned content agents call the model."""
    st = settings()
    store = MemoryStore()
    searched = []

    async def fake_search(query, domain, top):
        searched.append(query)
        return "[]"

    async def go():
        await store.create_run("r1", {"status": "queued"})
        cfg = RunConfig(rounds=2, stream=stream, max_output_tokens=500)
        return await execute(
            "r1",
            st,
            cfg,
            store,
            lambda rec: Model(st, "direct", rec, transport=asgi()),
            fake_search,
        )

    summary = run(go())
    assert store.runs["r1"]["status"] == "succeeded" and store.runs["r1"]["progress"] == "10/10"
    calls = store.calls["r1"]
    kinds = [c["kind"] for c in calls]
    # per round: planner (tool call + answer) and 3 assigned content agents (same); compute is skipped
    assert kinds.count("llm") == 2 * (2 + 3 * 2) and kinds.count("tool") == 2 * (1 + 3)
    assert {c["agent"] for c in calls if c["kind"] == "llm"} == {
        "planner",
        "cover",
        "network",
        "addressing",
    }
    assert summary["failed_calls"] == 0 and len(searched) == 8
    rnd = store.outputs["r1"]["rounds"][0]
    assert [s["sheet_name"] for s in rnd["plan"]["sections"]] == [
        "Cover",
        "Firewall rules",
        "Subnets",
    ]
    assert rnd["sections"]["compute"]["issues"] == ["no sections assigned"]
    assert all(len(rnd["sections"][t]["sheets"]) == 1 for t in ("cover", "network", "addressing"))
    assert set(store.files["r1"]) == {
        "lld-workbook.xlsx",
        "design-document.docx",
        "network-diagram.mmd",
    }


def test_failed_run_is_recorded():
    st = settings()
    store = MemoryStore()

    async def go():
        await store.create_run("r1", {"status": "queued"})
        make = lambda rec: Model(st, "direct", rec, transport=mock_transport(status=500))  # noqa: E731
        with pytest.raises(Exception):  # noqa: B017 - the agent framework wraps the HTTP error
            await execute("r1", st, RunConfig(max_output_tokens=100), store, make, None)

    run(go())
    assert store.runs["r1"]["status"] == "failed" and store.runs["r1"]["error"]
    assert any(c["http_status"] == 500 for c in store.calls["r1"])


def test_orchestration_is_a_maf_workflow():
    """planner -> fan-out -> four section executors -> fan-in -> assemble."""
    from worker.orchestrator import Deps, build_workflow

    d = Deps("r", RunConfig(agents=["network", "compute"]), [], None, None, None, MemoryStore(), 3)
    ids = {e.id for e in build_workflow(d).executors.values()}
    assert ids == {"planner", "network", "compute", "assemble"}


def test_assigned_sections_and_unknown_owner():
    from worker.orchestrator import assigned_sections

    plan = {"sections": [
        {"sheet_name": "A", "owner_agent": "network"},
        {"sheet_name": "B", "sheet_type": "subnets"},  # owner inferred from the type
        {"sheet_name": "C", "owner_agent": "nobody"},
        "junk",
    ]}  # fmt: skip
    assert [s["sheet_name"] for s in assigned_sections(plan, "network")] == ["A"]
    assert [s["sheet_name"] for s in assigned_sections(plan, "addressing")] == ["B"]
    assert assigned_sections({}, "cover") == []


def test_artifacts_open_and_survive_malformed_model_output(tmp_path):
    import io

    from docx import Document
    from openpyxl import load_workbook

    from worker.outputs import build_artifacts

    plan = {"engagement_id": "E", "sections": [{"sheet_name": "Cover", "order": 1, "owner_agent": "cover"}],
            "detected_vlans": [{"id": "10", "name": "Web", "cidr": "10.0.0.0/24", "platform": "onprem"}, "junk"]}  # fmt: skip
    sections = {
        "cover": {"sheets": [{"sheet_name": "Cover/1", "columns": ["a"], "rows": [["x"], "bad", [1, None]]}, "junk"]},
        "network": {"parse_error": "no JSON object in model output"},
        "compute": None,
    }  # fmt: skip
    files = build_artifacts(plan, sections)
    wb = load_workbook(io.BytesIO(files["lld-workbook.xlsx"]))
    assert wb.sheetnames == ["Plan", "Cover-1"] and wb["Cover-1"].max_row == 3
    texts = [p.text for p in Document(io.BytesIO(files["design-document.docx"])).paragraphs]
    assert "Cover/1" in texts and "Detected segments" in texts
    assert b"subgraph onprem" in files["network-diagram.mmd"]
    assert b'["Web<br/>VLAN 10<br/>10.0.0.0/24"]' in files["network-diagram.mmd"]
    assert b"No segments" in build_artifacts({}, {})["network-diagram.mmd"]


def test_matrix_items_are_finer_than_chunks():
    import json as _json
    import os

    from worker.prompts import build_prompt, to_items

    chunks = [_json.loads(line) for line in open(os.environ["FACTS_PATH"])]
    items = to_items(chunks)
    assert len(items) > 5 * len(chunks) and all({"id", "text"} == set(i) for i in items)
    system, user = build_prompt(7, items)
    assert _json.loads(user)["fact_count"] == 7 and "EXACTLY ONE" in system


def test_tracing_is_a_safe_noop_without_a_connection_string(monkeypatch):
    from worker import tracing

    monkeypatch.delenv("APPLICATIONINSIGHTS_CONNECTION_STRING", raising=False)
    assert tracing.setup_tracing("t") is False
    tracing.annotate_span({"ttfh_s": 1})
    tracing.flush_tracing()
    with tracing.run_span("x", a=1):
        pass


def test_long_answer_is_not_misreported_as_truncated():
    """The terminal event carries the whole answer and its marker leads it: it must still count."""
    big = {**fake_model._response("m", False, text="x" * 60_000)}
    ev = fake_model._sse({"type": "response.completed", "sequence_number": 3, "response": big})
    assert len(ev) > 60_000
    chunks = [ev[i : i + 8192] for i in range(0, len(ev), 8192)]
    res, _ = run(call(stream=True, transport=mock_transport(chunks=[event_chunk(), *chunks])))
    assert res.record["outcome"] == "completed" and res.record["completion_tokens"] == 7


def test_tracing_keeps_every_span(monkeypatch):
    """Azure Monitor's default sampler drops spans in bursts (observed): the lab must keep all of them."""
    import agent_framework.observability as maf_obs
    import azure.monitor.opentelemetry as azmon

    from worker import tracing

    seen = {}
    monkeypatch.setattr(azmon, "configure_azure_monitor", lambda **kw: seen.update(kw))
    monkeypatch.setattr(maf_obs, "enable_instrumentation", lambda **kw: None)
    monkeypatch.setattr(tracing, "_log_ended_spans", lambda: None)
    monkeypatch.setattr(tracing, "_enabled", False)
    monkeypatch.setenv("APPLICATIONINSIGHTS_CONNECTION_STRING", "InstrumentationKey=x")
    assert tracing.setup_tracing("t") is True
    assert seen["sampling_ratio"] == 1.0 and seen["connection_string"] == "InstrumentationKey=x"
    monkeypatch.setattr(tracing, "_enabled", False)


def test_agent_retries_with_backoff_and_records_each_retry():
    from worker.agents import with_retries

    calls, notes = {"n": 0}, []

    async def flaky():
        calls["n"] += 1
        if calls["n"] < 3:
            raise RuntimeError("exceeded token rate limit")
        return {"ok": True}

    async def note(attempt, err):
        notes.append((attempt, str(err)))

    out = run(with_retries(flaky, attempts=4, backoff_s=0.01, on_retry=note))
    assert out == {"ok": True} and calls["n"] == 3 and [a for a, _ in notes] == [1, 2]

    async def always():
        raise RuntimeError("down")

    with pytest.raises(RuntimeError):
        run(with_retries(always, attempts=2, backoff_s=0.01))
    with pytest.raises(RuntimeError):  # one attempt = no retry (the matrix setting)
        run(with_retries(always, attempts=1, backoff_s=0.01))


def test_pipeline_retries_a_failing_agent_then_succeeds():
    """First HTTP call fails (as a mid-stream rate limit would); the run still completes."""
    st = settings(llm_max_attempts=3, llm_retry_backoff_s=0.01)
    store = MemoryStore()
    state = {"n": 0}
    inner = asgi()

    class FlakyOnce(httpx.AsyncBaseTransport):
        async def handle_async_request(self, request):
            state["n"] += 1
            if state["n"] == 1:
                raise RuntimeError("rate limit inside the stream")
            return await inner.handle_async_request(request)

    async def search(q, d, t):
        return "[]"

    async def go():
        await store.create_run("r1", {"status": "queued"})
        cfg = RunConfig(max_output_tokens=500, stream=False)
        make = lambda rec: Model(st, "direct", rec, transport=FlakyOnce())  # noqa: E731
        await execute("r1", st, cfg, store, make, search)

    run(go())
    assert store.runs["r1"]["status"] == "succeeded"
    retries = [c for c in store.calls["r1"] if c["kind"] == "agent_retry"]
    assert len(retries) == 1 and retries[0]["agent"] == "planner" and retries[0]["error_type"]


def test_pacer_spaces_requests_to_the_token_budget():
    import time

    from worker.llm_client import TokenPacer

    async def go():
        pacer = TokenPacer(1000, window_s=0.4)
        t0 = time.monotonic()
        await pacer.acquire(600)  # fits
        await pacer.acquire(600)  # 1200 > 1000: waits for the first to age out of the window
        await pacer.acquire(5000)  # bigger than the whole budget: allowed once the window is empty
        return time.monotonic() - t0

    assert 0.7 <= run(go()) < 3.0
    assert TokenPacer.estimate(b"x" * 3500) == 1000


def test_rate_limit_inside_a_200_stream_is_labelled_not_called_a_cut():
    err_event = b'event: error\ndata: {"type":"error","message":"exceeded token rate limit"}\n\n'
    res, _ = run(call(stream=True, transport=mock_transport(chunks=[event_chunk(), err_event])))
    r = res.record
    assert (
        r["http_status"] == 200 and r["outcome"] == "truncated" and r["error_type"] == "RateLimit"
    )


def test_pacer_counts_the_whole_conversation_for_continuations():
    from worker.llm_client import TokenPacer

    pacer = TokenPacer(1000)
    chain = ("planner", 1)
    first = b'{"input": "' + b"x" * 3500 + b'"}'
    assert pacer.estimate_request(first, chain) == pacer.estimate(
        first
    )  # a fresh request: body size only
    pacer.observe(chain, 118_000)  # what the service reported for that turn (prompt + completion)
    cont = b'{"previous_response_id": "resp_1", "input": [{"type": "function_call_output"}]}'
    assert (
        pacer.estimate_request(cont, chain) >= 118_000
    )  # tiny body, but billed for the whole context
    assert pacer.estimate_request(cont, ("cover", 1)) < 100  # another agent's chain is unaffected
