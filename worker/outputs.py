"""Artifacts of a run: LLD workbook (.xlsx), design document (.docx) and a network diagram (Mermaid).

Built from the planner's section list and the content agents' sheets. Model output is untrusted, so every
value is coerced to text and malformed entries are skipped, never raised.
"""

import io
import re
from typing import Any

from docx import Document
from openpyxl import Workbook


def _txt(v: Any) -> str:
    return "" if v is None else str(v)


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", _txt(name)).strip()[:28] or "Sheet"
    title, n = base, 2
    while title.lower() in used:
        title, n = f"{base[:25]}-{n}", n + 1
    used.add(title.lower())
    return title


def _sheets(sections: dict[str, dict[str, Any]]) -> list[tuple[str, dict[str, Any]]]:
    out = []
    for owner, output in sections.items():
        for sheet in (output or {}).get("sheets") or []:
            if isinstance(sheet, dict):
                out.append((owner, sheet))
    return out


def _rows(sheet: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    cols = [_txt(c) for c in sheet.get("columns") or []]
    rows = [[_txt(c) for c in r] for r in sheet.get("rows") or [] if isinstance(r, list)]
    return cols, rows


def _plan_sections(plan: dict[str, Any]) -> list[dict[str, Any]]:
    secs = [s for s in plan.get("sections") or [] if isinstance(s, dict)]
    return sorted(
        secs, key=lambda s: s.get("order") if isinstance(s.get("order"), int | float) else 999
    )


def build_workbook(plan: dict[str, Any], sections: dict[str, dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws.append(["Order", "Sheet", "Type", "Owner", "Scope", "Rationale"])
    for s in _plan_sections(plan):
        ws.append(
            [
                _txt(s.get(k))
                for k in (
                    "order",
                    "sheet_name",
                    "sheet_type",
                    "owner_agent",
                    "scope_ref",
                    "rationale",
                )
            ]
        )
    used = {"plan"}
    for _, sheet in _sheets(sections):
        cols, rows = _rows(sheet)
        w = wb.create_sheet(_sheet_title(sheet.get("sheet_name") or sheet.get("section_id"), used))
        if cols:
            w.append(cols)
        for r in rows:
            w.append(r)
        if sheet.get("notes"):
            w.append([])
            w.append(["Notes", _txt(sheet["notes"])])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_document(plan: dict[str, Any], sections: dict[str, dict[str, Any]]) -> bytes:
    doc = Document()
    doc.add_heading("Low-level design (synthetic)", 0)
    doc.add_paragraph(
        f"Engagement: {_txt(plan.get('engagement_id'))}. Generated from a synthetic HLD."
    )
    doc.add_heading("Plan", 1)
    secs = _plan_sections(plan)
    if secs:
        t = doc.add_table(rows=1, cols=4)
        for i, h in enumerate(("Sheet", "Type", "Owner", "Rationale")):
            t.rows[0].cells[i].text = h
        for s in secs:
            cells = t.add_row().cells
            for i, k in enumerate(("sheet_name", "sheet_type", "owner_agent", "rationale")):
                cells[i].text = _txt(s.get(k))
    for title, key in (("Omitted candidates", "omitted_candidates"), ("Warnings", "warnings")):
        items = plan.get(key) or []
        if items:
            doc.add_heading(title, 2)
            for it in items:
                doc.add_paragraph(
                    _txt(it if not isinstance(it, dict) else " - ".join(map(_txt, it.values()))),
                    style="List Bullet",
                )
    vlans = [v for v in plan.get("detected_vlans") or [] if isinstance(v, dict)]
    if vlans:
        doc.add_heading("Detected segments", 1)
        t = doc.add_table(rows=1, cols=4)
        for i, h in enumerate(("VLAN", "Name", "CIDR", "Platform")):
            t.rows[0].cells[i].text = h
        for v in vlans:
            cells = t.add_row().cells
            for i, k in enumerate(("id", "name", "cidr", "platform")):
                cells[i].text = _txt(v.get(k))
    for owner, sheet in _sheets(sections):
        cols, rows = _rows(sheet)
        doc.add_heading(_txt(sheet.get("sheet_name") or sheet.get("section_id")), 1)
        doc.add_paragraph(f"Written by the {owner} agent.")
        if cols:
            t = doc.add_table(rows=1, cols=len(cols))
            for i, c in enumerate(cols):
                t.rows[0].cells[i].text = c
            for r in rows[:60]:  # keep the document readable; the workbook has every row
                cells = t.add_row().cells
                for i in range(len(cols)):
                    cells[i].text = r[i] if i < len(r) else ""
        if sheet.get("notes"):
            doc.add_paragraph(_txt(sheet["notes"]))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_diagram(plan: dict[str, Any]) -> str:
    """Mermaid flowchart of the detected segments, grouped by platform."""
    lines = ["flowchart LR"]
    groups: dict[str, list[str]] = {}
    for i, v in enumerate(v for v in plan.get("detected_vlans") or [] if isinstance(v, dict)):
        parts = [_txt(v.get("name")), f"VLAN {_txt(v.get('id'))}", _txt(v.get("cidr"))]
        label = "<br/>".join(re.sub(r'["<>\[\]{}|]', "", p) for p in parts if p.strip())
        platform = re.sub(r"\W", "_", _txt(v.get("platform"))) or "unknown"
        groups.setdefault(platform, []).append(f'    n{i}["{label}"]')
    for platform, nodes in groups.items():
        lines.append(f"  subgraph {platform}")
        lines += nodes
        lines.append("  end")
    if len(lines) == 1:
        lines.append('  none["No segments detected"]')
    return "\n".join(lines) + "\n"


def build_artifacts(plan: dict[str, Any], sections: dict[str, dict[str, Any]]) -> dict[str, bytes]:
    return {
        "lld-workbook.xlsx": build_workbook(plan, sections),
        "design-document.docx": build_document(plan, sections),
        "network-diagram.mmd": build_diagram(plan).encode(),
    }
